#!/usr/bin/env python3
"""Convert PowerPoint (.pptx) presentation to structured markdown text.

Extracts all slide titles, body paragraphs, bullet points, tables, hyperlinks,
and speaker notes without taking screenshots or requiring visual rendering.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, List

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    print("Error: python-pptx is required. Run: pip install python-pptx", file=sys.stderr)
    sys.exit(1)


class PptxTextExtractor:
    """Extracts text, tables, and notes from a PowerPoint presentation into Markdown.

    Attributes:
        pptx_path: Path to the input .pptx file.
        prs: Presentation object loaded via python-pptx.
    """

    def __init__(self, pptx_path: Path) -> None:
        """Initialize the extractor with the path to a PowerPoint presentation.

        Args:
            pptx_path: Absolute or relative path to the .pptx presentation.

        Raises:
            FileNotFoundError: If pptx_path does not exist.
        """
        if not pptx_path.exists():
            raise FileNotFoundError(f"Presentation file not found: {pptx_path}")
        self.pptx_path: Path = pptx_path
        self.prs: pptx.Presentation = pptx.Presentation(str(pptx_path))

    def extract_slide_to_markdown(self, slide: Any, slide_num: int) -> str:
        """Extract a single slide content into formatted Markdown.

        Args:
            slide: python-pptx Slide object.
            slide_num: 1-indexed slide number.

        Returns:
            Formatted markdown representation of the slide.
        """
        lines: List[str] = [f"## Slide {slide_num}"]

        # Process shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_frame = shape.text_frame
                for paragraph in text_frame.paragraphs:
                    p_text = self._format_paragraph(paragraph)
                    if p_text:
                        indent = "  " * paragraph.level
                        if paragraph.level > 0 or p_text.startswith("- ") or p_text.startswith("* "):
                            lines.append(f"{indent}* {p_text.lstrip('-* ')}")
                        else:
                            lines.append(f"{indent}{p_text}")

            elif shape.has_table:
                table = shape.table
                table_md = self._format_table(table)
                if table_md:
                    lines.append(table_md)

            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_desc = getattr(shape, "name", "Image")
                lines.append(f"*(Picture: {img_desc})*")

        # Process speaker notes if available
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_text = notes_tf.text.strip() if notes_tf else ""
            if notes_text:
                lines.append("\n> **Speaker Notes:**")
                for n_line in notes_text.splitlines():
                    if n_line.strip():
                        lines.append(f"> {n_line.strip()}")

        lines.append("\n---\n")
        return "\n".join(lines)

    def _format_paragraph(self, paragraph: Any) -> str:
        """Format paragraph runs with basic Markdown bold/italic and hyperlinks.

        Args:
            paragraph: python-pptx Paragraph object.

        Returns:
            Formatted string representation of paragraph text.
        """
        parts: List[str] = []
        for run in paragraph.runs:
            t = run.text
            if not t:
                continue
            if run.hyperlink and run.hyperlink.address:
                t = f"[{t}]({run.hyperlink.address})"
            if run.font.bold and run.font.italic:
                t = f"***{t}***"
            elif run.font.bold:
                t = f"**{t}**"
            elif run.font.italic:
                t = f"*{t}*"
            parts.append(t)

        full_text = "".join(parts).strip()
        return full_text

    def _format_table(self, table: Any) -> str:
        """Format a python-pptx Table into Markdown table syntax.

        Args:
            table: python-pptx Table object.

        Returns:
            Markdown table representation.
        """
        if not table.rows:
            return ""

        grid: List[List[str]] = []
        for row in table.rows:
            row_cells: List[str] = []
            for cell in row.cells:
                cell_text = " ".join(p.text.strip() for p in cell.text_frame.paragraphs if p.text.strip())
                row_cells.append(cell_text.replace("|", "\\|"))
            grid.append(row_cells)

        if not grid:
            return ""

        col_count = len(grid[0])
        md_rows: List[str] = []
        # Header
        md_rows.append("| " + " | ".join(grid[0]) + " |")
        md_rows.append("| " + " | ".join(["---"] * col_count) + " |")
        # Body
        for row in grid[1:]:
            md_rows.append("| " + " | ".join(row) + " |")

        return "\n" + "\n".join(md_rows) + "\n"

    def convert_to_markdown(self) -> str:
        """Convert the full presentation to a structured Markdown document.

        Returns:
            Complete Markdown text of the presentation.
        """
        title = self.pptx_path.stem
        header = f"# Presentation: {title}\n\n*Source: `{self.pptx_path.name}` ({len(self.prs.slides)} slides)*\n\n---\n\n"
        slides_md = [self.extract_slide_to_markdown(slide, idx + 1) for idx, slide in enumerate(self.prs.slides)]
        return header + "\n".join(slides_md)


def main() -> None:
    """CLI entry point for pptx_to_text."""
    if len(sys.argv) < 2:
        print("Usage: python pptx_to_text.py <path_to_presentation.pptx> [output.md]", file=sys.stderr)
        sys.exit(1)

    input_path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else None

    extractor = PptxTextExtractor(input_path)
    markdown_content = extractor.convert_to_markdown()

    if output_path:
        output_path.write_text(markdown_content, encoding="utf-8")
        print(f"Extracted {len(extractor.prs.slides)} slides to: {output_path}")
    else:
        print(markdown_content)


if __name__ == "__main__":
    main()
