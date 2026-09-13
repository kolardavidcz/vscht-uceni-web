#!/usr/bin/env python3
"""Audit all text boxes across all slides to ensure no leftover text."""

from pathlib import Path
import pptx

deck_path = Path("/mnt/c/Users/kolar/Downloads/strašení prváků - bioinformatika.pptx")
prs = pptx.Presentation(str(deck_path))

for i, slide in enumerate(prs.slides):
    slide_num = i + 1
    print(f"\n--- Slide {slide_num} ---")
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.strip():
            b = f"({s.left/914400:.1f}, {s.top/914400:.1f}, {s.width/914400:.1f}, {s.height/914400:.1f})"
            txt_preview = " | ".join(s.text_frame.text.strip().splitlines())[:80]
            print(f"  Shape '{s.name}' {b}: {txt_preview}")
