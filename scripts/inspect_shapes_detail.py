#!/usr/bin/env python3
"""Inspect detailed shape coordinates, texts, and images per slide."""

from pathlib import Path
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

deck_path = Path("/mnt/c/Users/kolar/Downloads/strašení prváků - bioinformatika.pptx")
prs = pptx.Presentation(str(deck_path))

for i, slide in enumerate(prs.slides):
    slide_num = i + 1
    print(f"\n==================== SLIDE {slide_num} ====================")
    for s in slide.shapes:
        left_in = s.left / 914400
        top_in = s.top / 914400
        w_in = s.width / 914400
        h_in = s.height / 914400
        
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"  [PICTURE] name='{s.name}', box=({left_in:.2f}, {top_in:.2f}, {w_in:.2f}, {h_in:.2f})")
        elif s.has_text_frame:
            full_txt = " \\n ".join(p.text.strip() for p in s.text_frame.paragraphs if p.text.strip())
            print(f"  [TEXT] name='{s.name}', box=({left_in:.2f}, {top_in:.2f}, {w_in:.2f}, {h_in:.2f}) -> {full_txt[:120]}")
        else:
            print(f"  [SHAPE] name='{s.name}', type={s.shape_type}, box=({left_in:.2f}, {top_in:.2f}, {w_in:.2f}, {h_in:.2f})")
