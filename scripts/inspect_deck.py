#!/usr/bin/env python3
"""Inspect presentation shapes, geometry, images, and text."""

from pathlib import Path
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

deck_path = Path("/mnt/c/Users/kolar/Downloads/strašení prváků - bioinformatika.pptx")
prs = pptx.Presentation(str(deck_path))

w_in = prs.slide_width / 914400
h_in = prs.slide_height / 914400
print(f"Presentation dimensions: {w_in:.3f}\" x {h_in:.3f}\" ({prs.slide_width} x {prs.slide_height} EMU)")
print(f"Total slides: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides):
    slide_num = i + 1
    shapes_info = []
    for s in slide.shapes:
        s_type = s.shape_type
        is_pic = (s_type == MSO_SHAPE_TYPE.PICTURE)
        text_preview = ""
        if s.has_text_frame:
            text_preview = " | ".join(p.text.strip() for p in s.text_frame.paragraphs if p.text.strip())[:60]
        shapes_info.append(f"[{s.name} (type={s_type}, pic={is_pic}): {text_preview}]")
    print(f"Slide {slide_num:2d}: {len(slide.shapes)} shapes -> {', '.join(shapes_info[:4])}")
